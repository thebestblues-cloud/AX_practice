"""대시보드 공헌이익·손익분기 Excel 검산(원자료 + SUMIFS 수식)"""
import pandas as pd, json
from openpyxl import Workbook
S=pd.read_csv('../data/참고서출고.csv',encoding='utf-8-sig');R=pd.read_csv('../data/반품입고.csv',encoding='utf-8-sig');M=pd.read_csv('../data/도서마스터.csv',encoding='utf-8-sig')
E=json.load(open('/tmp/lo/engine_out.json'))
wb=Workbook();ws=wb.active;ws.title='출고';ws.append(['월','거래처명','시리즈명','출고부수','공급단가','출고액','제작원가','출고원가'])
for i,r in enumerate(S.itertuples(index=False),start=2): ws.append(list(r)+[f'=D{i}*E{i}',f'=INDEX(마스터!$E:$E,MATCH(C{i},마스터!$A:$A,0))',f'=D{i}*G{i}'])
n=len(S)+1
wr=wb.create_sheet('반품');wr.append(['월','거래처명','시리즈명','반품부수','반품사유','단가','반품액','제작원가','재판매가능원가'])
for i,r in enumerate(R.itertuples(index=False),start=2):
    wr.append(list(r)+[f'=SUMIFS(출고!$F$2:$F${n},출고!$B$2:$B${n},B{i},출고!$C$2:$C${n},C{i})/SUMIFS(출고!$D$2:$D${n},출고!$B$2:$B${n},B{i},출고!$C$2:$C${n},C{i})',f'=D{i}*F{i}',f'=INDEX(마스터!$E:$E,MATCH(C{i},마스터!$A:$A,0))',f'=IF(OR(E{i}="개정판교체",E{i}="파본"),0,D{i}*H{i})'])
m=len(R)+1
wm=wb.create_sheet('마스터');wm.append(list(M.columns))
for r in M.fillna('').itertuples(index=False): wm.append(list(r))
v=wb.create_sheet('공헌이익',0)
v.append(['거래처','출고액','반품액','출고원가','재판매가능원가','반품부수','출고부수','재판매불가부수','공헌이익(수식)','공헌이익(대시보드)','차이','손익분기(수식)','손익분기(대시보드)','차이','물류비'])

for i,e in enumerate(E,start=2):
    p=e['n']
    v.append([p,f'=SUMIFS(출고!F:F,출고!B:B,A{i})',f'=SUMIFS(반품!G:G,반품!B:B,A{i})',f'=SUMIFS(출고!H:H,출고!B:B,A{i})',f'=SUMIFS(반품!I:I,반품!B:B,A{i})',f'=SUMIFS(반품!D:D,반품!B:B,A{i})',f'=SUMIFS(출고!D:D,출고!B:B,A{i})',
      f'=SUMIFS(반품!D:D,반품!B:B,A{i},반품!E:E,"개정판교체")+SUMIFS(반품!D:D,반품!B:B,A{i},반품!E:E,"파본")',
      f'=B{i}-C{i}-(D{i}-E{i})-F{i}*$O$2',e['contrib'],f'=I{i}-J{i}',
      f'=(B{i}/G{i}-D{i}/G{i})/(B{i}/G{i}-D{i}/G{i}+$O$2+H{i}/F{i}*D{i}/G{i})',e['be'],f'=L{i}-M{i}'])
v['O2']=1400
v.column_dimensions['A'].width=14
wb.save('검산_공헌이익.xlsx');print(len(E))
