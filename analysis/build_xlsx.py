"""Excel 검산 워크북: 원자료 + SUMIFS 수식 검산 시트 (LibreOffice 재계산 후 pandas 값과 대조)"""
import pandas as pd, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
o=json.load(open('analysis/results.json'))
S=pd.read_csv('data/참고서출고.csv',encoding='utf-8-sig'); R=pd.read_csv('data/반품입고.csv',encoding='utf-8-sig'); M=pd.read_csv('data/도서마스터.csv',encoding='utf-8-sig')
wb=Workbook(); ws=wb.active; ws.title='출고'
ws.append(list(S.columns)+['출고액','정가','정가액'])
for i,r in enumerate(S.itertuples(index=False),start=2):
    ws.append(list(r)+[f'=D{i}*E{i}',f'=INDEX(도서마스터!$D:$D,MATCH(C{i},도서마스터!$A:$A,0))',f'=D{i}*G{i}'])
n=len(S)+1
wr=wb.create_sheet('반품'); wr.append(list(R.columns)+['단가(거래처×시리즈 가중평균)','반품액','제작원가','폐기여부','폐기원가'])
for i,r in enumerate(R.itertuples(index=False),start=2):
    wr.append(list(r)+[f'=SUMIFS(출고!$F$2:$F${n},출고!$B$2:$B${n},B{i},출고!$C$2:$C${n},C{i})/SUMIFS(출고!$D$2:$D${n},출고!$B$2:$B${n},B{i},출고!$C$2:$C${n},C{i})',
      f'=D{i}*F{i}',f'=INDEX(도서마스터!$E:$E,MATCH(C{i},도서마스터!$A:$A,0))',f'=IF(OR(E{i}="개정판교체",E{i}="파본"),1,0)',f'=D{i}*H{i}*I{i}'])
m=len(R)+1
wm=wb.create_sheet('도서마스터'); wm.append(list(M.columns))
for r in M.itertuples(index=False): wm.append([None if (isinstance(x,float) and pd.isna(x)) else x for x in r])
v=wb.create_sheet('검산',0); bold=Font(bold=True)
v.append(['항목','Excel 수식값','Python 값','차이']); [setattr(c,'font',bold) for c in v[1]]
T=o['total']; X=o['synth']; h=o['high']; rv=o['revision']
rows=[('누적 출고부수',f'=SUM(출고!D2:D{n})',T['ship']),('누적 반품부수',f'=SUM(반품!D2:D{m})',T['ret']),
 ('출고 기준 매출',f'=SUM(출고!F2:F{n})',T['rev']),('반품액',f'=SUM(반품!G2:G{m})',T['retamt']),
 ('반품률',f'=B3/B2',T['rate']),('실판매 기준 매출',f'=B4-B5',T['netrev']),
 ('한솔북센 출고',f'=SUMIFS(출고!D2:D{n},출고!B2:B{n},"한솔북센")',1014080),
 ('한솔북센 반품',f'=SUMIFS(반품!D2:D{m},반품!B2:B{m},"한솔북센")',381708),
 ('한솔 반품률',f'=B9/B8',h['rate']),('자기 제외 평균 반품률',f'=(B3-B9)/(B2-B8)',h['ex_rate']),
 ('한솔 초과 반품부수',f'=B9-B8*B11',h['excess']),
 ('한솔 평균 공급단가',f'=SUMIFS(출고!F2:F{n},출고!B2:B{n},"한솔북센")/B8',h['unitprice']),
 ('초과반품 매출 연환산(24개월÷2)',f'=B12*B13/2',h['excess_amt_ann']),
 ('개정판교체 반품부수',f'=SUMIFS(반품!D2:D{m},반품!E2:E{m},"개정판교체")',rv['repl_total']),
 ('파본 반품부수',f'=SUMIFS(반품!D2:D{m},반품!E2:E{m},"파본")',o['reason']['pabon_qty']),
 ('폐기 제작원가(개정판교체+파본) 24개월',f'=SUM(반품!J2:J{m})',X['disp_24']),
 ('왕복 물류비 24개월(1,400원/부 가정)',f'=B3*1400',X['log_24']),
 ('실전만점 파본 부수',f'=SUMIFS(반품!D2:D{m},반품!E2:E{m},"파본",반품!C2:C{m},"실전만점*")',61114),
 ('실전만점 출고부수',f'=SUMIFS(출고!D2:D{n},출고!C2:C{n},"실전만점*")',1438010),
 ('파본 집중도(실전만점)',f'=(B19/B16)/(B20/B2)',[b for b in o['reason']['pabon_brand'] if b['브랜드']=='실전만점'][0]['conc']),
]
# 개정판 창 출고 (6종)
def add(mm,k): y,q=map(int,mm.split('-')); t=y*12+q-1+k; return f'{t//12}-{t%12+1:02d}'
for r in rv['rows']:
    w=[add(r['M'],k) for k in (-3,-2,-1)]; p=[add(r['M'],k) for k in (1,2,3)]
    f='+'.join([f'SUMIFS(출고!D2:D{n},출고!C2:C{n},"{r["시리즈"]}",출고!A2:A{n},"{x}")' for x in w])
    g='+'.join([f'SUMIFS(반품!D2:D{m},반품!C2:C{m},"{r["시리즈"]}",반품!A2:A{m},"{x}")' for x in p])
    rows.append((f'{r["시리즈"]} M-3~M-1 출고','='+f,r['win'])); rows.append((f'{r["시리즈"]} M+1~M+3 반품','='+g,r['post_ret']))
# 공급률
for k in ['북웨이브몰','페이지온몰','하루책방닷컴','한솔북센']:
    for y in ['2024','2025']:
        f=f'=SUMIFS(출고!F2:F{n},출고!B2:B{n},"{k}",출고!A2:A{n},"{y}*")/SUMIFS(출고!H2:H{n},출고!B2:B{n},"{k}",출고!A2:A{n},"{y}*")'
        val=[t for t in o['supply']['table'] if t['거래처명']==k][0][y]
        rows.append((f'{k} 공급률 {y}',f,val))
for i,(a,f,pv) in enumerate(rows,start=2):
    v.append([a,f,pv,f'=B{i}-C{i}'])
v.column_dimensions['A'].width=42; v.column_dimensions['B'].width=20; v.column_dimensions['C'].width=20
wb.save('analysis/검산_워크북.xlsx'); print(len(rows),'checks')
